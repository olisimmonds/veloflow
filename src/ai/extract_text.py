# import easyocr
from docx import Document
import PyPDF2
import pandas as pd
from bs4 import BeautifulSoup
import openpyxl
from io import BytesIO
import requests
import streamlit as st


def extract_text(file_path):
    file_type = file_path.split('.')[-1]
    print(f"{file_type=}")
    if file_type.startswith('pdf'):
        return extract_pdf_text(file_path)
    elif file_type.startswith('docx'):
        return extract_docx_text(file_path)
    elif file_type.startswith('txt'):
        return extract_txt_text(file_path)
    elif file_type.startswith('csv'):
        return extract_csv_text(file_path)
    # elif file_type.startswith('jpg'):
    #     return extract_image_text(file_path)
    # elif file_type.startswith('png'):
    #     return extract_image_text(file_path)
    elif file_type.startswith('html'):
        return extract_html_text(file_path)
    elif file_type.startswith('md'):
        return extract_markdown_text(file_path)
    elif file_type.startswith('xls'):
        return extract_excel_text(file_path)
    elif file_type.startswith('xlsx'):
        return extract_excel_text(file_path)
    else:
        raise ValueError("Unsupported file format")


def extract_pdf_text(url):
    # try:
    response = requests.get(url)
    response.raise_for_status()
    reader = PyPDF2.PdfReader(BytesIO(response.content))
    text = "".join(page.extract_text() for page in reader.pages if page.extract_text())
    return text
    # except Exception as e:
    #     st.write(f"Error extracting text from PDF: {e}")
    #     return ""


def extract_docx_text(url):
    response = requests.get(url)
    response.raise_for_status()
    doc = Document(BytesIO(response.content))
    return "\n".join(para.text for para in doc.paragraphs)


def extract_txt_text(url):
    response = requests.get(url)
    response.raise_for_status()
    return response.text


def extract_csv_text(url):
    response = requests.get(url)
    response.raise_for_status()
    df = pd.read_csv(BytesIO(response.content))
    return df.to_string()


# def extract_image_text(url):
#     response = requests.get(url)
#     response.raise_for_status()
#     reader = easyocr.Reader(['en'])  # Load the model for English
#     result = reader.readtext(BytesIO(response.content), detail=0)
#     return " ".join(result)


def extract_html_text(url):
    response = requests.get(url)
    response.raise_for_status()
    soup = BeautifulSoup(response.text, 'html.parser')
    return soup.get_text()


def extract_markdown_text(url):
    response = requests.get(url)
    response.raise_for_status()
    return response.text


def extract_excel_text(url):
    response = requests.get(url)
    response.raise_for_status()
    wb = openpyxl.load_workbook(BytesIO(response.content))
    sheet = wb.active
    return "\n".join(" ".join(map(str, row)) for row in sheet.iter_rows(values_only=True))
