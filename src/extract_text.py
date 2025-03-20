import easyocr
from docx import Document
import PyPDF2
import pandas as pd
from bs4 import BeautifulSoup
import openpyxl

def extract_text(file_path):
    if file_path.endswith('.pdf'):
        return extract_pdf_text(file_path)
    elif file_path.endswith('.docx'):
        return extract_docx_text(file_path)
    elif file_path.endswith('.txt'):
        return extract_txt_text(file_path)
    elif file_path.endswith('.csv'):
        return extract_csv_text(file_path)
    elif file_path.endswith(('.jpg', '.png')):
        return extract_image_text(file_path)
    elif file_path.endswith('.html'):
        return extract_html_text(file_path)
    elif file_path.endswith('.md'):
        return extract_markdown_text(file_path)
    elif file_path.endswith(('.xls', '.xlsx')):
        return extract_excel_text(file_path)
    else:
        raise ValueError("Unsupported file format")

def extract_pdf_text(file_path):
    with open(file_path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
    return text

def extract_docx_text(file_path):
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text
    return text

def extract_txt_text(file_path):
    with open(file_path, 'r') as file:
        text = file.read()
    return text

def extract_csv_text(file_path):
    df = pd.read_csv(file_path)
    return df.to_string()

def extract_image_text(file_path):
    reader = easyocr.Reader(['en'])  # Load the model for English
    result = reader.readtext(file_path, detail=0)
    return " ".join(result)

def extract_html_text(file_path):
    with open(file_path, 'r') as file:
        soup = BeautifulSoup(file, 'html.parser')
        text = soup.get_text()
    return text

def extract_markdown_text(file_path):
    with open(file_path, 'r') as file:
        text = file.read()
    return text

def extract_excel_text(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    text = ""
    for row in sheet.iter_rows(values_only=True):
        text += ' '.join(map(str, row)) + '\n'
    return text
